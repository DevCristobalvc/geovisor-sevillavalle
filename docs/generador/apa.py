# -*- coding: utf-8 -*-
"""Utilidades de formato APA 7 para python-docx.

Aplica: Times New Roman 12, interlineado doble, margenes de 1 pulgada,
numeracion de pagina en el encabezado, niveles de titulo APA, figuras
con rotulo y nota, tablas APA y referencias con sangria francesa.
"""
import struct
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"
SIZE = Pt(12)
CONTENT_WIDTH_IN = 6.5  # carta 8.5in - margenes de 1in


# ─────────────────────────── infraestructura ───────────────────────────

def png_size(path):
    with open(path, "rb") as f:
        head = f.read(24)
    return struct.unpack(">II", head[16:24])


def new_document():
    doc = Document()

    # Margenes de 1 pulgada (APA 7, 2.22)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Estilo base: Times New Roman 12, interlineado doble, sin espacio extra
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE
    normal.font.color.rgb = RGBColor(0, 0, 0)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.widow_control = True

    return doc


def add_page_numbers(doc):
    """Numero de pagina en la esquina superior derecha (APA 7, 2.18)."""
    for section in doc.sections:
        header_p = section.header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = header_p.add_run()
        run.font.name = FONT
        run.font.size = SIZE
        for instr, kind in (("begin", "fldChar"), ("PAGE", "instrText"), ("end", "fldChar")):
            el = OxmlElement("w:" + kind)
            if kind == "fldChar":
                el.set(qn("w:fldCharType"), instr)
            else:
                el.set(qn("xml:space"), "preserve")
                el.text = " PAGE "
            run._r.append(el)


def _style_run(run, bold=False, italic=False, size=SIZE):
    run.font.name = FONT
    run.font.size = size
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT)
    return run


def para(doc, text="", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False,
         indent_first=None, space_after=Pt(0), spacing=WD_LINE_SPACING.DOUBLE,
         keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = spacing
    pf.space_after = space_after
    pf.space_before = Pt(0)
    pf.keep_with_next = keep_with_next
    if indent_first is not None:
        pf.first_line_indent = Inches(indent_first)
    if text:
        _style_run(p.add_run(text), bold=bold, italic=italic)
    return p


def body(doc, text):
    """Parrafo de cuerpo: sangria de primera linea de 0.5in (APA 7, 2.24)."""
    return para(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_first=0.5)


def h1(doc, text):
    """Nivel 1: centrado, negrita, mayuscula inicial."""
    return para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, keep_with_next=True)


def h2(doc, text):
    """Nivel 2: alineado a la izquierda, negrita."""
    return para(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, keep_with_next=True)


def h3(doc, text):
    """Nivel 3: alineado a la izquierda, negrita y cursiva."""
    return para(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, italic=True,
                keep_with_next=True)


def bullet(doc, text, bold_lead=None):
    """Vineta con sangria; opcionalmente una etiqueta inicial en negrita."""
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.left_indent = Inches(0.75)
    if bold_lead:
        _style_run(p.add_run(bold_lead), bold=True)
    _style_run(p.add_run(text))
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.left_indent = Inches(0.75)
    _style_run(p.add_run(text))
    return p


# ─────────────────────────── figuras y tablas ───────────────────────────

def figure(doc, number, title, image_path, note=None, max_width_in=CONTENT_WIDTH_IN,
           max_height_in=6.5):
    """Figura en formato APA 7 (7.22).

    Rotulo "Figura N" en negrita, titulo en cursiva debajo, imagen centrada
    y nota explicativa opcional.
    """
    # Rotulo
    para(doc, f"Figura {number}", bold=True, space_after=Pt(0), keep_with_next=True)
    # Titulo en cursiva
    para(doc, title, italic=True, space_after=Pt(6), keep_with_next=True)

    # Escalado: encaja en el ancho util sin desbordar el alto de pagina
    w_px, h_px = png_size(image_path)
    width_in = max_width_in
    height_in = width_in * h_px / w_px
    if height_in > max_height_in:
        height_in = max_height_in
        width_in = height_in * w_px / h_px

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    # La nota debe quedarse con su imagen
    p.paragraph_format.keep_with_next = bool(note)
    p.add_run().add_picture(image_path, width=Inches(width_in))

    if note:
        np = doc.add_paragraph()
        np.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        np.paragraph_format.space_after = Pt(12)
        np.paragraph_format.space_before = Pt(0)
        _style_run(np.add_run("Nota. "), italic=True, size=Pt(10))
        _style_run(np.add_run(note), size=Pt(10))
    else:
        para(doc, "", space_after=Pt(6), spacing=WD_LINE_SPACING.SINGLE)


def _keep_row_together(row):
    """Evita que una fila se parta y la mantiene unida a lo que sigue.

    Sin esto, Word deja la nota de la tabla huerfana al comienzo de la pagina
    siguiente y la anterior queda medio vacia.
    """
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trPr.append(cant_split)
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.keep_with_next = True


def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom", "left", "right"):
        spec = kwargs.get(edge)
        el = OxmlElement(f"w:{edge}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec))
            el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def table(doc, number, title, headers, rows, note=None, col_widths=None):
    """Tabla en formato APA 7 (7.8): solo lineas horizontales."""
    para(doc, f"Tabla {number}", bold=True, space_after=Pt(0), keep_with_next=True)
    para(doc, title, italic=True, space_after=Pt(6), keep_with_next=True)

    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    # Encabezado: linea arriba y abajo
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.space_before = Pt(3)
        _style_run(p.add_run(htext), bold=True, size=Pt(11))
        _set_cell_border(cell, top=8, bottom=6)

    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        last = r_i == len(rows) - 1
        for i, val in enumerate(row):
            cell = cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            _style_run(p.add_run(str(val)), size=Pt(11))
            _set_cell_border(cell, bottom=8 if last else None)

    if col_widths:
        for row in t.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    # La tabla viaja como un bloque, arrastrando su nota
    for row in t.rows:
        _keep_row_together(row)

    if note:
        np = doc.add_paragraph()
        np.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        np.paragraph_format.space_after = Pt(12)
        np.paragraph_format.space_before = Pt(6)
        _style_run(np.add_run("Nota. "), italic=True, size=Pt(10))
        _style_run(np.add_run(note), size=Pt(10))
    else:
        para(doc, "", space_after=Pt(6), spacing=WD_LINE_SPACING.SINGLE)
    return t


def reference(doc, text):
    """Entrada de referencia con sangria francesa de 0.5in (APA 7, 9.43)."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.5)
    pf.space_after = Pt(0)
    _style_run(p.add_run(text))
    return p


def title_page(doc, title, subtitle, authors, affiliation, program, course,
               instructor, date):
    """Portada de trabajo academico APA 7 (2.3)."""
    for _ in range(4):
        para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)

    para(doc, title, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    if subtitle:
        para(doc, subtitle, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)

    for line in (authors, affiliation, program, course, instructor, date):
        para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_page_break()


def abstract(doc, text, keywords):
    """Resumen sin sangria y palabras clave en cursiva (APA 7, 2.9)."""
    h1(doc, "Resumen")
    para(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.space_after = Pt(0)
    _style_run(p.add_run("Palabras clave: "), italic=True)
    _style_run(p.add_run(keywords))
    doc.add_page_break()
