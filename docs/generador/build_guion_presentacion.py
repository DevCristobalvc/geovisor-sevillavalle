# -*- coding: utf-8 -*-
"""Genera el guion en espanol de la Presentacion 1: traduccion 1:1 del contenido de
cada diapositiva, separada por diapositiva. Sirve de apoyo para la exposicion oral.

Requisitos: pip install python-docx
Uso:       python build_guion_presentacion.py "../../Presentacion 1 - Guion en espanol.docx"
"""
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

SP = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(SP, "Guion Presentacion 1.docx")

GREEN = RGBColor(0x2D, 0x6A, 0x4F)
GRAY = RGBColor(0x5F, 0x6B, 0x7A)

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin = s.bottom_margin = Cm(2.5)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)


def slide(n, title, kicker=None):
    if n > 1:
        doc.add_page_break()
    p = doc.add_paragraph()
    r = p.add_run(f"DIAPOSITIVA {n}")
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = GREEN
    if kicker:
        r2 = p.add_run(f"   ·   {kicker}")
        r2.font.size = Pt(9)
        r2.font.color.rgb = GRAY
    h = doc.add_paragraph()
    hr = h.add_run(title)
    hr.font.size = Pt(18)
    hr.font.bold = True
    h.paragraph_format.space_after = Pt(8)


def label(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = GRAY
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)


def text(t, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.italic = italic
    p.paragraph_format.space_after = Pt(6)
    return p


def bullet(t):
    p = doc.add_paragraph(t, style="List Bullet")
    p.paragraph_format.space_after = Pt(4)


def numbered(t):
    p = doc.add_paragraph(t, style="List Number")
    p.paragraph_format.space_after = Pt(4)


def card(title, body):
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.font.bold = True
    p.paragraph_format.space_after = Pt(0)
    text(body)


# ───────────────────────────── Portada del documento ─────────────────────────────
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Presentación 1 — Guion en español")
r.font.size = Pt(22)
r.font.bold = True
s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run("Traducción 1:1 del contenido de cada diapositiva\n"
              "Desarrollo de un Geovisor Web Ecopedagógico para Sevilla, Valle del Cauca")
r.font.size = Pt(12)
r.font.color.rgb = GRAY
n = doc.add_paragraph()
n.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = n.add_run("Las diapositivas están en inglés según el criterio del corte 1; la "
              "exposición oral se hace en español. Este documento reproduce, "
              "diapositiva por diapositiva, exactamente lo que aparece en pantalla.")
r.font.size = Pt(10)
r.italic = True
r.font.color.rgb = GRAY

# ───────────────────────────── 1. Portada ─────────────────────────────
slide(1, "Portada")
text("UNIVERSIDAD SANTIAGO DE CALI")
text("Facultad de Ingeniería · Programa de Ingeniería de Sistemas")
text("PROPUESTA DE TRABAJO DE GRADO · PRESENTACIÓN 1")
label("Título")
text("Desarrollo de un Geovisor Web Ecopedagógico para Sevilla, Valle del Cauca")
label("Subtítulo")
text("Un WebGIS que conecta datos geoespaciales oficiales con contenido pedagógico para "
     "fomentar el reconocimiento crítico de las transformaciones socioecosistémicas "
     "entre estudiantes de educación básica secundaria.")
label("Autores")
text("Cristóbal Valencia Cerón\nJosé David Molina Delgado")
label("Directores de trabajo de grado")
text("Diego Fernando Loaiza\nSilvia Andrea Quijano Pérez")
text("Grupo de Investigación COMBA I+D · Línea: Desarrollo de Sistemas Informáticos")
text("Santiago de Cali, septiembre de 2026")
label("Recuadro del QR")
text("PLATAFORMA EN LÍNEA\nwww.refiup.app\nEscanea para abrir el Geovisor")

# ───────────────────────────── 2. Agenda ─────────────────────────────
slide(2, "Agenda", "Presentación 1 · 8 minutos")
for i, it in enumerate([
    "Introducción (contexto)", "Planteamiento del problema", "Pregunta de investigación",
    "Justificación", "Marco teórico", "Estado del arte", "Objetivo general",
    "Objetivos específicos", "Metodología y avance", "Referencias"]):
    text(f"{i + 1:02d}  {it}")

# ───────────────────────────── 3. Introducción ─────────────────────────────
slide(3, "Introducción", "Contexto")
bullet("Sevilla, norte del Valle del Cauca, Colombia: un municipio reconocido por su "
       "vocación cafetera y su rico patrimonio biocultural, parte del Paisaje Cultural "
       "Cafetero de la UNESCO (2011).")
bullet("En las últimas décadas, las transformaciones socioecosistémicas —cambios en el "
       "uso del suelo, la crisis cafetera y el cambio climático— han generado una "
       "creciente desconexión entre las comunidades locales y su territorio.")
bullet("La educación territorial emerge como una estrategia clave para fortalecer la "
       "conciencia ambiental y la identidad local.")
bullet("Este proyecto propone un Geovisor web con enfoque ecopedagógico: cartografía "
       "interactiva, datos geoespaciales oficiales y contenido pedagógico multimedia "
       "para estudiantes de básica secundaria.")
label("Tarjeta: Articulación")
text("Desarrollado en articulación con la tesis Cartografiar las huellas del café "
     "(J. Rodríguez Camacho), que construye el contenido pedagógico.")
text("Este proyecto se enfoca en la plataforma web: interfaces de usuario, el modelo de "
     "datos geoespacial y la especificación técnica.")
text("Referente técnico y funcional: Portal GeoCVC (geo.cvc.gov.co), Corporación "
     "Autónoma Regional del Valle del Cauca.")

# ───────────────────────────── 4. Problema ─────────────────────────────
slide(4, "Planteamiento del problema", "Problema")
card("Sin herramienta integrada",
     "Sevilla no cuenta con una herramienta digital que muestre, de manera integrada y "
     "pedagógica, sus transformaciones socioecosistémicas: actores sociales, agua, "
     "biodiversidad y dinámicas de uso del suelo.")
card("Brecha tecnológica en el aula",
     "Los docentes de básica secundaria carecen de herramientas interactivas para "
     "contextualizar los contenidos curriculares con la realidad geográfica y "
     "socioambiental local.")
card("Datos oficiales dispersos",
     "La información geoespacial del IGAC, la CVC, el IDEAM y el DANE está dispersa en "
     "múltiples plataformas y formatos técnicos no accesibles para usuarios no "
     "especializados.")
card("Exigencia ecopedagógica",
     "La conciencia ecológica requiere estrategias didácticas que articulen la "
     "experiencia territorial con el pensamiento crítico y sistémico (Zimmermann, 2005).")
label("Banda inferior")
text("Problema central: la ausencia de una herramienta que integre cartografía, "
     "multimedia y pedagogía para el contexto específico de Sevilla.")

# ───────────────────────────── 5. Pregunta ─────────────────────────────
slide(5, "Pregunta de investigación", "Pregunta")
text("¿Cuáles son los requerimientos técnicos, arquitecturales y de experiencia de "
     "usuario que debe cumplir un Geovisor web ecopedagógico para promover el "
     "reconocimiento crítico de las transformaciones socioecosistémicas en Sevilla, "
     "Valle del Cauca?")

# ───────────────────────────── 6. Justificación ─────────────────────────────
slide(6, "Justificación", "Por qué este proyecto")
card("Tecnológica",
     "Los SIG web han democratizado el acceso a la cartografía digital (p. ej., GeoCVC); "
     "sin embargo, no están diseñados con enfoque pedagógico ni adaptados a usuarios de "
     "básica secundaria.\n\nUn sistema construido sobre estándares OGC (WMS), GeoJSON "
     "estático y tecnologías web modernas (React, TypeScript, Leaflet) es una "
     "contribución concreta a los SIG educativos.")
card("Educativa",
     "La ecopedagogía busca superar la educación ambiental convencional para lograr una "
     "transformación real de las actitudes frente al territorio.\n\nEl Geovisor ofrece "
     "fichas pedagógicas, multimedia y recorridos guiados alineados con los lineamientos "
     "del Ministerio de Educación Nacional (MEN) para Ciencias Naturales y Educación "
     "Ambiental.")
card("Territorial",
     "Sevilla enfrenta transformaciones urgentes: expansión ganadera en páramos, "
     "conflictos de uso del suelo, deterioro de cuencas y pérdida de la identidad "
     "cafetera.\n\nHacerlas visibles mediante capas geográficas conectadas a narrativas "
     "pedagógicas fortalece el tejido social y la conciencia ambiental local.")
label("Nota al pie")
text("Relevancia académica: una metodología de diseño de WebGIS ecopedagógicos, "
     "replicable en otros municipios del Valle del Cauca y de Colombia.", italic=True)

# ───────────────────────────── 7. Marco teórico ─────────────────────────────
slide(7, "Marco teórico", "Fundamentos")
card("Ecopedagogía y educación territorial",
     "Zimmermann (2005), Ecopedagogía: el planeta en emergencia. Una relación "
     "hombre-ambiente desde una perspectiva psicoambiental, orientada a una conciencia "
     "ecológica que se exprese en el comportamiento cotidiano. En este proyecto guía las "
     "decisiones de diseño: la selección de capas como narrativa territorial, la "
     "progresividad pedagógica en la interfaz y los módulos de reflexión crítica "
     "ligados al mapa.")
card("Sistemas de Información Geográfica web (WebGIS)",
     "Plataformas para visualizar, consultar y analizar datos geoespaciales desde un "
     "navegador, sin software especializado (Esri España, 2018). Interoperabilidad "
     "mediante OGC WMS 1.3.0 (OGC, 2010). Fundamentos en Maguire, Goodchild y Rhind "
     "(1991); ecosistema SIG de código abierto en Steiniger y Hunter (2013).")
card("Territorio y socioecosistemas en Sevilla",
     "Parte del Paisaje Cultural Cafetero de la UNESCO (2011). Páramo Chilí-Barragán, "
     "un sector del PNN Las Hermosas (65,58 ha), bosque andino, humedales y bosque seco "
     "tropical —del cual solo sobrevive el 1,7 % de la cobertura original en el Valle "
     "del Cauca (INCIVA, 2023)—. Cuencas: ríos Bugalagrande y La Vieja.")
card("Marco legal",
     "Ley 99 de 1993 (Sistema Nacional Ambiental, SINA) · Ley 115 de 1994 (Ley General "
     "de Educación: la educación ambiental como área obligatoria) · Ley 1523 de 2012 "
     "(gestión del riesgo de desastres) · Ley 1581 de 2012 (protección de datos "
     "personales) · Lineamientos de la ICDE y el IGAC para datos geoespaciales "
     "oficiales.")

# ───────────────────────────── 8. Estado del arte ─────────────────────────────
slide(8, "Estado del arte", "Plataformas e investigaciones relacionadas")
label("Tabla comparativa")
tbl = doc.add_table(rows=1, cols=4)
tbl.style = "Light Grid Accent 1"
hdr = tbl.rows[0].cells
for i, h in enumerate(["Plataforma / trabajo", "Alcance", "Datos oficiales",
                       "Enfoque pedagógico"]):
    hdr[i].text = h
rows = [
    ("Portal GeoCVC (CVC, 2024)", "Geoportal ambiental regional, Valle del Cauca",
     "Sí — WMS/WFS", "No"),
    ("Geoservicios IGAC (IGAC, 2024)", "Cartografía básica y catastral nacional",
     "Sí — WMS", "No"),
    ("OpenStreetMap (Haklay y Weber, 2008)",
     "Mapas callejeros colaborativos, generados por usuarios", "Comunidad", "No"),
    ("Ecosistema SIG libre (Steiniger y Hunter, 2013)",
     "Herramientas de código abierto: Leaflet, PostGIS, GeoServer…", "Herramientas", "No"),
    ("Geovisor Ecopedagógico de Sevilla (este proyecto)",
     "WebGIS municipal con fichas, glosario y recorridos guiados",
     "Sí — CVC, IGAC, IDEAM, Humboldt, RUNAP", "Sí"),
]
for row in rows:
    cells = tbl.add_row().cells
    for i, v in enumerate(row):
        cells[i].text = v
doc.add_paragraph()
label("Nota al pie")
text("Brecha identificada: no existe un WebGIS para Sevilla que combine capas "
     "geoespaciales oficiales, fichas ecopedagógicas y recorridos guiados diseñados "
     "para estudiantes de básica secundaria. Los portales existentes atienden a usuarios "
     "técnicos.")

# ───────────────────────────── 9. Objetivo general ─────────────────────────────
slide(9, "Objetivo general", "Objetivos")
text("Desarrollar un Geovisor web con enfoque ecopedagógico que promueva el "
     "reconocimiento crítico de las transformaciones socioecosistémicas en Sevilla "
     "(Valle del Cauca), contribuyendo al fortalecimiento de la conciencia territorial "
     "y la resignificación del patrimonio biocultural.")

# ───────────────────────────── 10. Objetivos específicos ─────────────────────────────
slide(10, "Objetivos específicos", "Objetivos")
numbered("Arquitectura y modelo de datos — Especificar la arquitectura de software del "
         "Geovisor (React/TypeScript, Leaflet.js, servicios WMS, GeoJSON estático) y el "
         "modelo de datos geoespacial para sus seis categorías temáticas.")
numbered("Interfaces y requisitos — Definir los requisitos funcionales y no funcionales "
         "y diseñar las interfaces de usuario (wireframes y prototipos) con usabilidad, "
         "accesibilidad y coherencia ecopedagógica.")
numbered("Integración de datos — Integrar fuentes geoespaciales oficiales (CVC, IGAC, "
         "IDEAM, Instituto Humboldt, RUNAP) mediante servicios WMS públicos y archivos "
         "GeoJSON.")

# ───────────────────────────── 11. Metodología ─────────────────────────────
slide(11, "Metodología", "Cuatro fases")
card("01 — Revisión bibliográfica y análisis de requerimientos",
     "WebGIS, ecopedagogía y sistemas de información territorial en educación. Inventario "
     "de fuentes geoespaciales oficiales para Sevilla. Requisitos funcionales y no "
     "funcionales, en articulación con la tesis colaboradora.")
card("02 — Diseño del sistema",
     "Arquitectura de software y stack tecnológico. Modelo de datos geoespacial con seis "
     "categorías temáticas. Wireframes de las pantallas principales.")
card("03 — Implementación y pruebas",
     "Aplicación web estática que integra capas oficiales. Pruebas de usabilidad con "
     "estudiantes de básica secundaria.")
card("04 — Documentación y cierre",
     "Documentación técnica (manuales de usuario y técnico) y documento final de la "
     "tesis.")
label("Nota al pie")
text("Instrumentos: matrices de requerimientos, diagramas de arquitectura, wireframes, "
     "registros de pruebas de usabilidad.   Indicadores: arquitectura documentada; modelo "
     "de datos con seis categorías temáticas; prototipo funcional desplegado; "
     "especificación de requerimientos validada.")

# ───────────────────────────── 12. Avance ─────────────────────────────
slide(12, "Avance actual", "El Geovisor hoy")
label("Cifras")
text("23 capas temáticas · 6 categorías · 23 fichas pedagógicas · 4 recorridos guiados · "
     "5 proveedores WMS oficiales")
bullet("Aplicación de página única que se ejecuta íntegramente en el navegador: React 18, "
       "TypeScript, Vite, Tailwind CSS y Leaflet — sin servidor, sin base de datos, sin "
       "datos personales.")
bullet("Capas de la CVC, el IGAC, el IDEAM, el Instituto Humboldt y el RUNAP mediante OGC "
       "WMS, más GeoJSON versionado para vectores pequeños.")
bullet("Aplicación web progresiva: funciona sin conexión después de la primera visita — "
       "diseñada para aulas rurales con conectividad intermitente.")
bullet("Fichas pedagógicas con preguntas de reflexión, glosario, herramientas de "
       "medición, exportación de imagen y enlaces de mapa compartibles.")
bullet("Desplegado públicamente en www.refiup.app con despliegue continuo.")
label("Recuadro del QR")
text("www.refiup.app")

# ───────────────────────────── 13. Referencias ─────────────────────────────
slide(13, "Referencias", "APA 7.ª edición")
text("Las referencias se conservan tal como aparecen en la diapositiva (formato APA; los "
     "títulos en español no se traducen y los títulos en inglés se mantienen en su idioma "
     "original).", italic=True)
for r in [
    "Congreso de Colombia. (1993). Ley 99 de 1993, por la cual se crea el Ministerio del "
    "Medio Ambiente y se organiza el Sistema Nacional Ambiental (SINA). Diario Oficial "
    "No. 41.146.",
    "Congreso de Colombia. (1994). Ley 115 de 1994, Ley General de Educación. Diario "
    "Oficial No. 41.214.",
    "Congreso de Colombia. (2012a). Ley 1523 de 2012, por la cual se adopta la política "
    "nacional de gestión del riesgo de desastres. Diario Oficial No. 48.411.",
    "Congreso de Colombia. (2012b). Ley 1581 de 2012, por la cual se dictan disposiciones "
    "generales para la protección de datos personales. Diario Oficial No. 48.587.",
    "Corporación Autónoma Regional del Valle del Cauca. (2024). Portal GeoCVC. "
    "https://geo.cvc.gov.co/inicio/",
    "Esri España. (2018). ¿Qué es un WebGIS? https://www.esri.es",
    "Haklay, M., & Weber, P. (2008). OpenStreetMap: User-generated street maps. IEEE "
    "Pervasive Computing, 7(4), 12–18. https://doi.org/10.1109/MPRV.2008.80",
    "Instituto Geográfico Agustín Codazzi. (2024). Geoservicios cartográficos de "
    "Colombia. https://www.igac.gov.co",
    "Instituto para la Investigación y la Preservación del Patrimonio Cultural y Natural "
    "del Valle del Cauca. (2023). Bosque seco tropical en el Valle del Cauca. INCIVA.",
    "Maguire, D. J., Goodchild, M. F., & Rhind, D. W. (1991). Geographical information "
    "systems: Principles and applications. Longman Scientific & Technical.",
    "Open Geospatial Consortium. (2010). OpenGIS Web Map Service (WMS) implementation "
    "specification (Version 1.3.0). https://www.ogc.org/standards/wms",
    "Parques Nacionales Naturales de Colombia. (s.f.). Parque Nacional Natural Las "
    "Hermosas. https://www.parquesnacionales.gov.co",
    "Steiniger, S., & Hunter, A. J. S. (2013). The 2012 free and open source GIS software "
    "map: A guide to facilitate research, development, and adoption. Computers, "
    "Environment and Urban Systems, 39, 136–150. "
    "https://doi.org/10.1016/j.compenvurbsys.2012.10.003",
    "UNESCO. (2011). Paisaje Cultural Cafetero de Colombia. Lista del Patrimonio Mundial. "
    "https://whc.unesco.org/en/list/1121",
    "Zimmermann, M. (2005). Ecopedagogía: El planeta en emergencia. Ecoe Ediciones.",
]:
    p = text(r)
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)

# ───────────────────────────── 14. Cierre ─────────────────────────────
slide(14, "Cierre")
text("Gracias")
text("Preguntas y discusión")
label("Contacto")
text("cristobal.valencia00@usc.edu.co\njose.molina03@usc.edu.co")
text("Universidad Santiago de Cali · Ingeniería de Sistemas · COMBA I+D")
label("Recuadro del QR")
text("www.refiup.app\nAbrir el Geovisor")

doc.save(OUT)
print("Generado:", OUT)
